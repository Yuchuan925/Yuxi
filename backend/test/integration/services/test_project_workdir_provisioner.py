"""真实 Sandbox 的 UserWorkspace 挂载契约测试。"""

from __future__ import annotations

import asyncio
import shutil
import uuid

import pytest

from yuxi.modules.agents.runtime.sandbox import ProvisionerSandboxBackend, get_sandbox_provider
from yuxi.modules.workspace.paths import (
    ensure_user_workspace,
    global_user_data_dir,
    user_workspace_dir,
    workspace_uid_dirname,
)
from yuxi.modules.extensions.skills.projection import get_user_skills_root_dir, sync_user_accessible_skills
from yuxi.infrastructure.runtime_settings import get_skill_projection_dir, get_user_data_dir

pytestmark = [pytest.mark.asyncio, pytest.mark.integration]


def _create_workdir(uid: str):
    """在真实 UserWorkspace 中创建测试 Workdir。"""
    ensure_user_workspace(uid)
    get_user_skills_root_dir(uid).mkdir(parents=True, exist_ok=True)
    workdir_id = str(uuid.uuid4())
    workdir_path = f"projects/{workdir_id}"
    host_workdir = user_workspace_dir(uid) / "projects" / workdir_id
    host_workdir.mkdir(parents=True)
    return workdir_path, host_workdir


def _cleanup_user_storage(uid: str) -> None:
    """清理真实 provisioner 测试创建的 uid 级持久目录。"""
    shutil.rmtree(global_user_data_dir(uid), ignore_errors=True)
    shutil.rmtree(get_skill_projection_dir() / workspace_uid_dirname(uid), ignore_errors=True)


async def test_ephemeral_remote_skill_sandbox_does_not_create_persistent_uid_roots():
    """无环境的一次性 Sandbox 只能使用 emptyDir/tmpfs，不得消耗持久卷 inode。"""
    suffix = uuid.uuid4().hex
    uid = f"remote-skill-{suffix}"
    scope = uid
    safe_uid = workspace_uid_dirname(uid)
    user_root = get_user_data_dir() / "shared" / safe_uid
    skill_root = get_skill_projection_dir() / safe_uid
    backend = ProvisionerSandboxBackend(
        thread_id=scope,
        uid=uid,
        inherit_env=False,
    )
    try:
        result = await asyncio.to_thread(
            backend.execute,
            "mkdir -p /home/gem/user-data/outputs && printf ephemeral > /home/gem/user-data/outputs/check.txt",
        )
        assert result.exit_code == 0, result.output
        assert not user_root.exists()
        assert not skill_root.exists()
    finally:
        try:
            await asyncio.to_thread(
                get_sandbox_provider().release,
                scope,
                uid=uid,
                clear_cache_on_delete_failure=True,
            )
        finally:
            shutil.rmtree(user_root, ignore_errors=True)
            shutil.rmtree(skill_root, ignore_errors=True)


async def test_two_sandboxes_share_project_files_but_not_runtime_state():
    suffix = uuid.uuid4().hex
    uid = f"pytest-project-{suffix}"
    workdir_path, _ = _create_workdir(uid)
    first_scope = f"pytest-runtime-a-{suffix}"
    second_scope = f"pytest-runtime-b-{suffix}"
    project_root = f"/home/gem/user-data/{workdir_path}"
    project_file = f"{project_root}/outputs/shared.txt"
    runtime_file = f"/tmp/yuxi-runtime-{suffix}"

    first = ProvisionerSandboxBackend(thread_id=first_scope, uid=uid, workdir_path=workdir_path)
    second = ProvisionerSandboxBackend(thread_id=second_scope, uid=uid, workdir_path=workdir_path)

    try:
        first_result = await asyncio.to_thread(
            first.execute,
            f"mkdir -p {project_root}/outputs && printf shared-bytes > {project_file} && printf private-runtime > {runtime_file}",
        )
        assert first_result.exit_code == 0, first_result.output

        read_result = await asyncio.to_thread(second.execute, f"cat {project_file}")
        assert read_result.exit_code == 0, read_result.output
        assert read_result.output == "shared-bytes"

        runtime_result = await asyncio.to_thread(second.execute, f"test ! -e {runtime_file}")
        assert runtime_result.exit_code == 0, runtime_result.output

        first_connection = get_sandbox_provider().get(
            first_scope,
            uid=uid,
            workdir_path=workdir_path,
        )
        second_connection = get_sandbox_provider().get(
            second_scope,
            uid=uid,
            workdir_path=workdir_path,
        )
        assert first_connection is not None and second_connection is not None
        assert first_connection.sandbox_id != second_connection.sandbox_id
        assert first_connection.generation and second_connection.generation
        assert first_connection.generation != second_connection.generation
    finally:
        try:
            await asyncio.to_thread(first.execute, f"rm -f {project_file} {runtime_file}")
        except Exception:
            pass
        for scope in (first_scope, second_scope):
            try:
                await asyncio.to_thread(
                    get_sandbox_provider().release,
                    scope,
                    uid=uid,
                    workdir_path=workdir_path,
                    clear_cache_on_delete_failure=True,
                )
            except Exception:
                pass
        _cleanup_user_storage(uid)


async def test_recreated_runtime_keeps_project_files_and_drops_process_state():
    suffix = uuid.uuid4().hex
    uid = f"pytest-project-recreate-{suffix}"
    workdir_path, _ = _create_workdir(uid)
    first_scope = f"pytest-runtime-before-{suffix}"
    second_scope = f"pytest-runtime-after-{suffix}"
    project_root = f"/home/gem/user-data/{workdir_path}"
    project_file = f"{project_root}/outputs/persistent.txt"
    runtime_file = f"/tmp/yuxi-runtime-{suffix}"
    provider = get_sandbox_provider()

    first = ProvisionerSandboxBackend(thread_id=first_scope, uid=uid, workdir_path=workdir_path)
    second = ProvisionerSandboxBackend(thread_id=second_scope, uid=uid, workdir_path=workdir_path)
    try:
        result = await asyncio.to_thread(
            first.execute,
            f"mkdir -p {project_root}/outputs && printf persistent > {project_file} && printf transient > {runtime_file}",
        )
        assert result.exit_code == 0, result.output
        await asyncio.to_thread(
            provider.release,
            first_scope,
            uid=uid,
            workdir_path=workdir_path,
            clear_cache_on_delete_failure=True,
        )

        result = await asyncio.to_thread(second.execute, f"cat {project_file} && test ! -e {runtime_file}")
        assert result.exit_code == 0, result.output
        assert result.output == "persistent"
    finally:
        try:
            await asyncio.to_thread(second.execute, f"rm -f {project_file}")
        except Exception:
            pass
        for scope in (first_scope, second_scope):
            try:
                await asyncio.to_thread(
                    provider.release,
                    scope,
                    uid=uid,
                    workdir_path=workdir_path,
                    clear_cache_on_delete_failure=True,
                )
            except Exception:
                pass
        _cleanup_user_storage(uid)


async def test_workspace_file_remains_available_when_execution_runtime_is_released():
    """执行 runtime 被删除后，UserWorkspace 中的文件仍由宿主持有。"""
    suffix = uuid.uuid4().hex
    uid = f"pytest-file-bridge-{suffix}"
    workdir_path, host_workdir = _create_workdir(uid)
    runtime_scope = f"pytest-runtime-{suffix}"
    project_root = f"/home/gem/user-data/{workdir_path}"
    project_file = f"{project_root}/outputs/realtime.txt"
    provider = get_sandbox_provider()
    runtime_backend = ProvisionerSandboxBackend(
        thread_id=runtime_scope,
        uid=uid,
        workdir_path=workdir_path,
    )
    try:
        result = await asyncio.to_thread(
            runtime_backend.execute,
            f"mkdir -p {project_root}/outputs && printf realtime > {project_file}",
        )
        assert result.exit_code == 0, result.output
        await asyncio.to_thread(
            provider.release,
            runtime_scope,
            uid=uid,
            workdir_path=workdir_path,
            clear_cache_on_delete_failure=True,
        )

        assert (host_workdir / "outputs" / "realtime.txt").read_text(encoding="utf-8") == "realtime"
    finally:
        try:
            await asyncio.to_thread(
                provider.release,
                runtime_scope,
                uid=uid,
                workdir_path=workdir_path,
                clear_cache_on_delete_failure=True,
            )
        except Exception:
            pass
        _cleanup_user_storage(uid)


async def test_user_skill_projection_is_shared_across_sandboxes_but_isolated_by_uid(tmp_path):
    """同一用户的 Sandbox 共享授权 Skill 文件，不同用户不可读。"""
    suffix = uuid.uuid4().hex
    uid = f"pytest-skills-{suffix}"
    other_uid = f"pytest-skills-other-{suffix}"
    first_scope = f"pytest-skills-runtime-a-{suffix}"
    second_scope = f"pytest-skills-runtime-b-{suffix}"
    other_scope = f"pytest-skills-runtime-other-{suffix}"
    selected_source = tmp_path / "selected"
    unselected_source = tmp_path / "authorized-unselected"
    selected_source.mkdir()
    unselected_source.mkdir()
    (selected_source / "SKILL.md").write_text("selected-skill", encoding="utf-8")
    (unselected_source / "SKILL.md").write_text("authorized-unselected-skill", encoding="utf-8")

    await asyncio.to_thread(
        sync_user_accessible_skills,
        uid,
        {
            "selected": selected_source,
            "authorized-unselected": unselected_source,
        },
    )
    await asyncio.to_thread(sync_user_accessible_skills, other_uid, {})
    ensure_user_workspace(uid)
    ensure_user_workspace(other_uid)

    first = ProvisionerSandboxBackend(thread_id=first_scope, uid=uid)
    second = ProvisionerSandboxBackend(thread_id=second_scope, uid=uid)
    other = ProvisionerSandboxBackend(thread_id=other_scope, uid=other_uid)
    try:
        write_result = await asyncio.to_thread(
            first.execute,
            "printf tampered > /home/gem/skills/selected/SKILL.md",
        )
        first_selected = await asyncio.to_thread(first.read, "/home/gem/skills/selected/SKILL.md")
        second_unselected = await asyncio.to_thread(
            second.read,
            "/home/gem/skills/authorized-unselected/SKILL.md",
        )
        other_selected = await asyncio.to_thread(other.read, "/home/gem/skills/selected/SKILL.md")

        assert write_result.exit_code != 0
        assert first_selected.error is None
        assert first_selected.file_data == {"content": "selected-skill", "encoding": "utf-8"}
        assert second_unselected.error is None
        assert second_unselected.file_data == {
            "content": "authorized-unselected-skill",
            "encoding": "utf-8",
        }
        assert other_selected.file_data is None
        assert other_selected.error
        other_error = other_selected.error.lower()
        selected_skill_path = "/home/gem/skills/selected/skill.md"
        assert selected_skill_path in other_error
        canonical_not_found = f"file '{selected_skill_path}' not found"
        assert any(marker in other_error for marker in ("does not exist", canonical_not_found, "filenotfounderror"))
    finally:
        await asyncio.to_thread(sync_user_accessible_skills, uid, {})
        await asyncio.to_thread(sync_user_accessible_skills, other_uid, {})
        for scope, scope_uid in (
            (first_scope, uid),
            (second_scope, uid),
            (other_scope, other_uid),
        ):
            try:
                await asyncio.to_thread(
                    get_sandbox_provider().release,
                    scope,
                    uid=scope_uid,
                    clear_cache_on_delete_failure=True,
                )
            except Exception:
                pass
        _cleanup_user_storage(uid)
        _cleanup_user_storage(other_uid)


async def test_ocr_directory_copy_is_visible_in_real_sandbox(tmp_path):
    """真实 OCR 工具复制 Office 目录，沙盒可直接读取 Markdown 和图片。"""
    import io
    from pathlib import PurePosixPath
    from types import SimpleNamespace

    from docx import Document
    from PIL import Image

    from yuxi.bootstrap.models import load_models
    from yuxi.infrastructure.postgres.manager import pg_manager
    from yuxi.modules.agents.runtime.sandbox.paths import runtime_workdir_path
    from yuxi.modules.extensions.tools.builtin.ocr_parse_file import ocr_parse_file

    load_models()
    suffix = uuid.uuid4().hex
    uid = f"pytest-ocr-copy-{suffix}"
    runtime_scope = f"pytest-ocr-copy-{suffix}"
    workdir_path, host_workdir = _create_workdir(uid)
    backend = ProvisionerSandboxBackend(thread_id=runtime_scope, uid=uid, workdir_path=workdir_path)
    try:
        picture = tmp_path / "chart.png"
        Image.new("RGB", (20, 20), "red").save(picture)
        document = Document()
        document.add_paragraph("directory copy fixture")
        document.add_picture(str(picture))
        document.save(host_workdir / "source.docx")
        runtime_path = runtime_workdir_path(workdir_path)
        context = {
            "runtime_scope_id": runtime_scope,
            "uid": uid,
            "workdir_relative_path": workdir_path,
            "workdir_path": runtime_path,
        }
        runtime = SimpleNamespace(config={"configurable": context}, context=SimpleNamespace(**context), state={})
        result = await ocr_parse_file.coroutine(file_path=f"{runtime_path}/source.docx", runtime=runtime, ocr_engine="disable")
        [markdown_response] = await asyncio.to_thread(backend.download_files, [result["parsed_path"]])
        assert markdown_response.error is None
        markdown = markdown_response.content.decode()
        assert "directory copy fixture" in markdown
        assert "![image](images/figure-001.png)" in markdown
        image_path = str(PurePosixPath(result["parsed_path"]).parent / "images/figure-001.png")
        [image_response] = await asyncio.to_thread(backend.download_files, [image_path])
        assert image_response.error is None
        with Image.open(io.BytesIO(image_response.content)) as image:
            assert image.size == (20, 20)
            assert image.convert("RGB").getpixel((0, 0)) == (255, 0, 0)
        host_output = host_workdir / result["parsed_path"].removeprefix(runtime_path + "/")
        assert host_output.read_bytes() == markdown_response.content
        assert host_output.with_name("images").joinpath("figure-001.png").read_bytes() == image_response.content
    finally:
        await asyncio.to_thread(
            get_sandbox_provider().release,
            runtime_scope,
            uid=uid,
            workdir_path=workdir_path,
            clear_cache_on_delete_failure=True,
        )
        _cleanup_user_storage(uid)
        await pg_manager.close()
